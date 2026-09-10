"""Purpose: extracts readable text from document formats (SPEC-01 req 2) as ordered units a
chunker can cite - a heading section, a slide, a whole document. The subtle failure here is not
a crash but plausible-looking output: python-docx exposes paragraphs and tables as separate
collections so body order must be recovered from the XML, and an HTML page's script and style
bodies read as text unless they are removed. Tables are counted as detected and as extracted
separately, because req 2 forbids reporting the first as if it were the second.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from ca_agent.config.settings import TextSettings
from ca_agent.core.enums import ErrorCategory, FormatFamily, UnitType
from ca_agent.core.model import ErrorInfo
from ca_agent.core.text import TextUnit

_STAGE = "text_extraction"
_HEADING_STYLE = re.compile(r"^heading\s*\d*$", re.IGNORECASE)
_TRAILING_SPACE = re.compile(r"[ \t]+$", re.MULTILINE)
_REPEATED_SPACE = re.compile(r"[ \t]{2,}")
_REPEATED_BLANK_LINE = re.compile(r"\n{3,}")
#: Slide part names are ppt/slides/slide12.xml; ordering must be numeric, not lexicographic.
_SLIDE_PART = re.compile(r"^ppt/slides/slide(\d+)\.xml$")
_DRAWING_TEXT = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"

_TEXT_FAMILIES = {
    FormatFamily.WORD_OOXML,
    FormatFamily.PRESENTATION_OOXML,
    FormatFamily.RTF,
    FormatFamily.HTML,
    FormatFamily.PLAIN_TEXT,
    FormatFamily.EMAIL,
}

_EMAIL_HEADERS = ("From", "To", "Cc", "Subject", "Date")


@dataclass(frozen=True, slots=True)
class TextResult:
    """Everything one document produced. A failure is returned, never raised.

    ``tables_detected`` and ``tables_extracted`` are deliberately separate counters: SPEC-01
    req 2 requires the format document to distinguish content that was found from content that
    was successfully read.
    """

    reader: str
    units: tuple[TextUnit, ...] = ()
    tables_detected: int = 0
    tables_extracted: int = 0
    encoding: str | None = None
    failure: ErrorInfo | None = None
    warnings: tuple[ErrorInfo, ...] = ()

    def character_count(self) -> int:
        return sum(len(unit.text) for unit in self.units)

    def has_content(self) -> bool:
        return any(unit.has_content() for unit in self.units)


def read_text(source: Path, *, settings: TextSettings, family: FormatFamily) -> TextResult:
    """Extract text from one document as ordered units."""
    if family not in _TEXT_FAMILIES:
        raise ValueError(f"{family.value} is not a text family; reader selection is wrong")

    try:
        payload = source.read_bytes()
    except OSError as error:
        return TextResult(
            reader="filesystem",
            failure=ErrorInfo(
                category=ErrorCategory.FILE_ACCESS_ERROR, message=str(error), stage=_STAGE
            ),
        )

    try:
        if family is FormatFamily.WORD_OOXML:
            return _read_docx(payload, settings)
        if family is FormatFamily.PRESENTATION_OOXML:
            return _read_pptx(payload, settings)
        if family is FormatFamily.HTML:
            return _read_html(payload, settings)
        if family is FormatFamily.EMAIL:
            return _read_email(payload, settings)
        if family is FormatFamily.RTF:
            return _read_rtf(payload, settings)
        return _read_plain(payload, settings)
    except _ReaderFailure as failure:
        return TextResult(reader=failure.reader, failure=failure.as_error())


class _ReaderFailure(Exception):
    """A whole-document failure. Problems with one element are recorded as warnings instead."""

    def __init__(self, reader: str, category: ErrorCategory, message: str) -> None:
        super().__init__(message)
        self.reader = reader
        self.category = category

    def as_error(self) -> ErrorInfo:
        return ErrorInfo(
            category=self.category, message=str(self), stage=_STAGE, reader=self.reader
        )


# --- word documents ----------------------------------------------------------------------


def _read_docx(payload: bytes, settings: TextSettings) -> TextResult:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    reader = "python-docx"
    try:
        document = docx.Document(io.BytesIO(payload))
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
        raise _ReaderFailure(reader, ErrorCategory.CORRUPT_FILE, str(error)) from error

    builder = _UnitBuilder(settings)
    warnings: list[ErrorInfo] = []
    detected = extracted = 0
    paragraph_tag, table_tag = qn("w:p"), qn("w:tbl")

    # The body is walked directly because document.paragraphs and document.tables are two
    # separate sequences: reading them in turn silently moves every table to the end.
    for element in document.element.body.iterchildren():
        if element.tag == paragraph_tag:
            paragraph = Paragraph(element, document)
            if _is_heading(paragraph):
                builder.start_heading(paragraph.text)
            else:
                builder.add(paragraph.text)
        elif element.tag == table_tag:
            detected += 1
            try:
                builder.add(_docx_table_text(Table(element, document)))
            except (ValueError, KeyError, IndexError, AttributeError) as error:
                warnings.append(
                    ErrorInfo(
                        category=ErrorCategory.EXTRACTION_ERROR,
                        message=f"table {detected} could not be read: {error}",
                        stage=_STAGE,
                        reader=reader,
                    )
                )
            else:
                extracted += 1

    return TextResult(
        reader=reader,
        units=builder.build(),
        tables_detected=detected,
        tables_extracted=extracted,
        warnings=tuple(warnings),
    )


def _is_heading(paragraph) -> bool:
    style = getattr(paragraph.style, "name", "") or ""
    return bool(_HEADING_STYLE.match(style)) and bool(paragraph.text.strip())


def _docx_table_text(table) -> str:
    """Render a table as pipe-separated rows so column association survives as text."""
    return "\n".join(
        " | ".join(cell.text.strip() for cell in row.cells) for row in table.rows
    )


# --- presentations -------------------------------------------------------------------------


def _read_pptx(payload: bytes, settings: TextSettings) -> TextResult:
    """Read slide text straight from the package XML.

    The corpus holds two presentations, which does not justify adding python-pptx; lxml is
    already a dependency and a slide's text is just its drawingml <a:t> runs in document order.
    """
    from lxml import etree

    reader = "lxml-pptx"
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as package:
            slides = sorted(
                (
                    (int(match.group(1)), name)
                    for name in package.namelist()
                    if (match := _SLIDE_PART.match(name))
                ),
            )
            units: list[TextUnit] = []
            for sequence, (number, name) in enumerate(slides):
                try:
                    root = etree.fromstring(package.read(name))
                except etree.XMLSyntaxError:
                    continue
                text = "\n".join(
                    node.text for node in root.iter(_DRAWING_TEXT) if node.text
                )
                units.append(
                    TextUnit(
                        unit_type=UnitType.PAGE,
                        unit_ref=f"slide:{number}",
                        text=_normalise(text, settings),
                        sequence=sequence,
                    )
                )
    except (zipfile.BadZipFile, KeyError, OSError) as error:
        raise _ReaderFailure(reader, ErrorCategory.CORRUPT_FILE, str(error)) from error

    return TextResult(reader=reader, units=tuple(units))


# --- html --------------------------------------------------------------------------------


def _read_html(payload: bytes, settings: TextSettings) -> TextResult:
    reader = "lxml.html"
    text, detected, extracted = _html_to_text(payload, settings, reader)
    return TextResult(
        reader=reader,
        units=_single_unit(text, "document"),
        tables_detected=detected,
        tables_extracted=extracted,
    )


def _html_to_text(payload: bytes, settings: TextSettings, reader: str) -> tuple[str, int, int]:
    """Reduce HTML to text, returning the table detected and extracted counts alongside."""
    import lxml.html
    from lxml import etree

    try:
        tree = lxml.html.document_fromstring(payload)
    except (etree.ParserError, etree.XMLSyntaxError, ValueError) as error:
        raise _ReaderFailure(reader, ErrorCategory.CORRUPT_FILE, str(error)) from error

    # Script and style bodies are markup machinery, not content; with_tail=False keeps the
    # text that follows the element, which is ordinary page content.
    etree.strip_elements(tree, "script", "style", with_tail=False)

    tables = tree.findall(".//table")
    extracted = 0
    for table in tables:
        try:
            table.text_content()
        except (ValueError, UnicodeDecodeError):
            # Drop it so the reported text and the extracted count cannot disagree.
            table.getparent().remove(table)
        else:
            extracted += 1

    return _normalise(tree.text_content(), settings), len(tables), extracted


# --- email --------------------------------------------------------------------------------


def _read_email(payload: bytes, settings: TextSettings) -> TextResult:
    import email
    from email import policy

    reader = "email"
    try:
        message = email.message_from_bytes(payload, policy=policy.default)
    except (ValueError, IndexError) as error:
        raise _ReaderFailure(reader, ErrorCategory.CORRUPT_FILE, str(error)) from error

    units: list[TextUnit] = []
    headers = "\n".join(
        f"{name}: {message[name]}" for name in _EMAIL_HEADERS if message[name] is not None
    )
    if headers:
        units.append(
            TextUnit(
                unit_type=UnitType.DOCUMENT,
                unit_ref="headers",
                text=_normalise(headers, settings),
                sequence=0,
            )
        )

    warnings = tuple(
        ErrorInfo(
            category=ErrorCategory.NO_COMPATIBLE_READER,
            message=f"attachment {part.get_filename() or 'unnamed'} was not extracted",
            stage=_STAGE,
            reader=reader,
        )
        for part in message.iter_attachments()
    )

    body = message.get_body(preferencelist=("plain", "html"))
    if body is not None:
        units.append(
            TextUnit(
                unit_type=UnitType.DOCUMENT,
                unit_ref="body",
                text=_email_body_text(body, settings, reader),
                sequence=len(units),
            )
        )

    return TextResult(reader=reader, units=tuple(units), warnings=warnings)


def _email_body_text(body, settings: TextSettings, reader: str) -> str:
    """Take the plain-text body, or reduce an HTML-only body to text."""
    content = body.get_content()
    if body.get_content_subtype() == "html":
        text, _, _ = _html_to_text(content.encode("utf-8", "replace"), settings, reader)
        return text
    return _normalise(content, settings)


# --- rtf and plain text -----------------------------------------------------------------------


def _read_rtf(payload: bytes, settings: TextSettings) -> TextResult:
    from striprtf.striprtf import rtf_to_text

    reader = "striprtf"
    decoded, encoding = _decode(payload, settings)
    try:
        text = rtf_to_text(decoded, errors="ignore")
    except (ValueError, IndexError, KeyError) as error:
        raise _ReaderFailure(reader, ErrorCategory.EXTRACTION_ERROR, str(error)) from error

    return TextResult(
        reader=reader,
        units=_single_unit(_normalise(text, settings), "document"),
        encoding=encoding,
    )


def _read_plain(payload: bytes, settings: TextSettings) -> TextResult:
    reader = "charset-normalizer"
    decoded, encoding = _decode(payload, settings)
    return TextResult(
        reader=reader,
        units=_single_unit(_normalise(decoded, settings), "document"),
        encoding=encoding,
    )


def _decode(payload: bytes, settings: TextSettings) -> tuple[str, str | None]:
    """Decode bytes to text, reporting the encoding used.

    Detection is tried first because the corpus mixes UTF-8, UTF-16 and cp1252 and carries
    Devanagari client names; the configured fallbacks only run when detection abstains.
    """
    from charset_normalizer import from_bytes

    best = from_bytes(payload).best()
    if best is not None:
        return str(best), best.encoding

    for candidate in settings.encoding_fallbacks:
        try:
            return payload.decode(candidate), candidate
        except (UnicodeDecodeError, LookupError):
            continue
    # latin-1 maps every byte, so this cannot lose the file entirely; the loss is fidelity.
    return payload.decode("latin-1", errors="replace"), None


# --- shared unit construction --------------------------------------------------------------------


class _UnitBuilder:
    """Accumulates document blocks into units, splitting at each heading.

    Content before the first heading becomes one document-level unit, so a document with no
    headings at all still yields its text in body order rather than nothing.
    """

    def __init__(self, settings: TextSettings) -> None:
        self._settings = settings
        self._units: list[TextUnit] = []
        self._blocks: list[str] = []
        self._heading: str | None = None

    def start_heading(self, heading: str) -> None:
        self._flush()
        self._heading = heading.strip()
        self._blocks = [self._heading]

    def add(self, block: str) -> None:
        if block.strip():
            self._blocks.append(block)

    def build(self) -> tuple[TextUnit, ...]:
        self._flush()
        return tuple(self._units)

    def _flush(self) -> None:
        text = _normalise("\n".join(self._blocks), self._settings)
        if text.strip():
            self._units.append(
                TextUnit(
                    unit_type=UnitType.HEADING if self._heading else UnitType.DOCUMENT,
                    unit_ref=self._heading or "document",
                    text=text,
                    sequence=len(self._units),
                )
            )
        self._blocks = []
        self._heading = None


def _single_unit(text: str, unit_ref: str) -> tuple[TextUnit, ...]:
    return (
        TextUnit(unit_type=UnitType.DOCUMENT, unit_ref=unit_ref, text=text, sequence=0),
    )


def _normalise(text: str, settings: TextSettings) -> str:
    """Collapse layout whitespace without destroying paragraph structure.

    Blank lines are meaningful in a statement or a note, so they are preserved up to two; runs
    of spaces from table layout and trailing whitespace are not.
    """
    if not settings.normalise_whitespace:
        return text
    collapsed = _REPEATED_SPACE.sub(" ", text.replace("\r\n", "\n").replace("\r", "\n"))
    collapsed = _TRAILING_SPACE.sub("", collapsed)
    return _REPEATED_BLANK_LINE.sub("\n\n", collapsed).strip()
